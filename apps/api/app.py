from flask import Flask, request, jsonify, send_from_directory, Response
from flask_cors import CORS
import os, io, json, hashlib, uuid
from datetime import datetime, date
from dateutil import parser as dtparser
from PIL import Image
try:
    # HEIC/HEIF support
    from pillow_heif import register_heif_opener
    register_heif_opener()
except Exception:
    pass

import openpyxl
from openpyxl.utils import get_column_letter

# --- support relative (package) & direct run ---
try:
    from . import config
    from .db import get_conn, init_db, init_dirs, now
except ImportError:
    import config
    from db import get_conn, init_db, init_dirs, now

app = Flask(__name__)
CORS(app, supports_credentials=True)
init_db()  # safe if already created

# --- Auth helpers ---
def make_token():
    raw = (config.ADMIN_PASSWORD + config.SECRET_SALT).encode("utf-8")
    return hashlib.sha256(raw).hexdigest()

ADMIN_TOKEN = make_token()

def require_admin():
    auth = request.headers.get("Authorization", "")
    if auth.startswith("Bearer "):
        token = auth.split(" ", 1)[1].strip()
        return token == ADMIN_TOKEN
    return False

# --- Utilities ---
ALLOWED_STATUSES = ["Not Started", "In Progress", "On Hold", "Done"]
ALLOWED_CATEGORIES = ["Operational", "Development", "Experiment/Belajar"]

def valid_transition(old, new):
    allowed = {
        "Not Started": {"In Progress", "On Hold", "Done"},
        "In Progress": {"On Hold", "Done"},
        "On Hold": {"In Progress", "Done"},
        "Done": {"In Progress"},  # Reopen -> In Progress
    }
    return new in allowed.get(old, set())

def save_image_webp(file_storage):
    """Terima file gambar ekstensi apa pun, simpan selalu .webp di /uploads/YYYY/MM/uuid.webp"""
    ts = datetime.utcnow()
    subdir = os.path.join(config.UPLOAD_DIR, str(ts.year), f"{ts.month:02d}")
    os.makedirs(subdir, exist_ok=True)
    name = f"{uuid.uuid4().hex}.webp"
    full_path = os.path.join(subdir, name)

    # Baca gambar (dukung HEIC via pillow-heif)
    try:
        img = Image.open(file_storage.stream)
    except Exception:
        file_storage.stream.seek(0)
        img = Image.open(io.BytesIO(file_storage.read()))

    # Pastikan RGB sebelum simpan ke webp
    if img.mode not in ("RGB", "L"):
        img = img.convert("RGB")
    else:
        img = img.convert("RGB")

    # Guard ukuran
    w, h = img.size
    m = max(w, h)
    if m > 2000:
        s = 2000 / m
        img = img.resize((int(w * s), int(h * s)), Image.LANCZOS)

    img.save(full_path, "WEBP", quality=90, method=6)

    # Kembalikan URL absolut yang konsisten untuk FE
    return f"/uploads/{ts.year}/{ts.month:02d}/{name}"

def task_to_dict(row, conn):
    d = dict(row)
    # Attach tags
    cur = conn.cursor()
    cur.execute(
        """
        SELECT t.name FROM tags t
        JOIN task_tags tt ON tt.tag_id = t.id
        WHERE tt.task_id = ?
        ORDER BY t.name ASC
        """,
        (row["id"],)
    )
    d["tags"] = [r["name"] for r in cur.fetchall()]
    # stale hint
    stale = False
    if d.get("status") == "In Progress" and d.get("last_update_at"):
        try:
            last = dtparser.isoparse(d["last_update_at"])
            delta = datetime.utcnow() - last.replace(tzinfo=None)
            stale = delta.days >= 3
        except Exception:
            stale = False
    d["stale"] = stale
    # deadline obj
    d["deadline"] = {
        "type": d.pop("deadline_type", None),
        "date": d.pop("deadline_date", None),
    }
    return d

def parse_pagination():
    page = max(1, int(request.args.get("page", 1)))
    size = min(100, max(1, int(request.args.get("size", 20))))
    offset = (page - 1) * size
    return page, size, offset

def apply_task_filters(base_sql, params):
    status = request.args.get("status")
    category = request.args.get("category")
    tag = request.args.get("tag")
    q = request.args.get("q")
    wheres = []
    if status in ALLOWED_STATUSES:
        wheres.append("tasks.status = ?")
        params.append(status)
    if category in ALLOWED_CATEGORIES:
        wheres.append("tasks.category = ?")
        params.append(category)
    if tag:
        wheres.append("""
            tasks.id IN (
              SELECT tt.task_id FROM task_tags tt
              JOIN tags tg ON tg.id = tt.tag_id
              WHERE tg.name = LOWER(?)
            )
        """)
        params.append(tag)
    if q:
        wheres.append("(tasks.title LIKE ?)")
        params.append(f"%{q}%")
    if wheres:
        base_sql += " WHERE " + " AND ".join(wheres)
    return base_sql, params

# --- Public: serve uploads (harus tepat ke folder uploads) ---
@app.route("/uploads/<path:subpath>")
def serve_upload(subpath):
    return send_from_directory(config.UPLOAD_DIR, subpath)

# --- Auth ---
@app.post("/api/login")
def login():
    data = request.get_json(force=True)
    username = data.get("username", "")
    password = data.get("password", "")
    if username == config.ADMIN_USERNAME and password == config.ADMIN_PASSWORD:
        return jsonify({"ok": True, "token": ADMIN_TOKEN, "username": username})
    return jsonify({"ok": False, "error": "Invalid credentials"}), 401

@app.get("/api/me")
def me():
    return jsonify({"is_admin": require_admin(), "username": config.ADMIN_USERNAME})

# --- Tags suggest ---
@app.get("/api/tags")
def tags_suggest():
    q = (request.args.get("suggest") or "").strip().lower()
    conn = get_conn()
    cur = conn.cursor()
    if q:
        cur.execute("SELECT name FROM tags WHERE name LIKE ? ORDER BY name LIMIT 20", (f"{q}%",))
    else:
        cur.execute("SELECT name FROM tags ORDER BY name LIMIT 50")
    names = [r["name"] for r in cur.fetchall()]
    conn.close()
    return jsonify({"items": names})

# --- Counts ---
@app.get("/api/tasks_counts")
def tasks_counts():
    conn = get_conn()
    cur = conn.cursor()
    cur.execute("SELECT status, COUNT(*) c FROM tasks GROUP BY status")
    counts = {r["status"]: r["c"] for r in cur.fetchall()}
    for s in ALLOWED_STATUSES:
        counts.setdefault(s, 0)
    conn.close()
    return jsonify(counts)

# --- List tasks ---
@app.get("/api/tasks")
def list_tasks():
    page, size, offset = parse_pagination()
    conn = get_conn()
    cur = conn.cursor()

    sql = "SELECT tasks.* FROM tasks"
    params = []
    sql, params = apply_task_filters(sql, params)

    ordering = """
      ORDER BY CASE tasks.status
        WHEN 'Not Started' THEN 1
        WHEN 'In Progress' THEN 2
        WHEN 'On Hold' THEN 3
        WHEN 'Done' THEN 4
      END ASC, tasks.updated_at DESC, tasks.id DESC
      LIMIT ? OFFSET ?
    """
    cur.execute("SELECT COUNT(1) AS total FROM (" + sql + ")", params)
    total = cur.fetchone()["total"]

    cur.execute(sql + ordering, params + [size, offset])
    rows = cur.fetchall()
    items = [task_to_dict(r, conn) for r in rows]
    conn.close()
    return jsonify({"page": page, "size": size, "total": total, "items": items})

@app.get("/api/tasks/<int:task_id>")
def get_task(task_id):
    conn = get_conn()
    cur = conn.cursor()
    cur.execute("SELECT * FROM tasks WHERE id = ?", (task_id,))
    row = cur.fetchone()
    if not row:
        conn.close()
        return jsonify({"error": "Not found"}), 404
    task = task_to_dict(row, conn)
    cur.execute("SELECT * FROM updates WHERE task_id = ? ORDER BY created_at DESC, id DESC", (task_id,))
    updates = [dict(r) for r in cur.fetchall()]
    task["updates"] = updates
    conn.close()
    return jsonify(task)

@app.post("/api/tasks")
def create_task():
    if not require_admin():
        return jsonify({"error": "Unauthorized"}), 401
    data = request.get_json(force=True)
    title = (data.get("title") or "").strip()
    category = data.get("category")
    tags = data.get("tags") or []

    if not title:
        return jsonify({"error": "Title required"}), 400
    if category not in ALLOWED_CATEGORIES:
        return jsonify({"error": "Invalid category"}), 400

    conn = get_conn()
    cur = conn.cursor()
    ts = now()
    cur.execute("""
      INSERT INTO tasks(title, category, status, created_at, updated_at)
      VALUES(?, ?, 'Not Started', ?, ?)
    """, (title, category, ts, ts))
    task_id = cur.lastrowid

    for t in tags:
        tag = (t or "").strip().lower()
        if not tag: continue
        try:
            cur.execute("INSERT INTO tags(name) VALUES(?)", (tag,))
        except Exception:
            pass
        cur.execute("SELECT id FROM tags WHERE name = ?", (tag,))
        tag_id = cur.fetchone()["id"]
        cur.execute("INSERT OR IGNORE INTO task_tags(task_id, tag_id) VALUES(?,?)", (task_id, tag_id))

    conn.commit()
    cur.execute("SELECT * FROM tasks WHERE id = ?", (task_id,))
    task = task_to_dict(cur.fetchone(), conn)
    conn.close()
    return jsonify(task), 201

@app.patch("/api/tasks/<int:task_id>")
def update_task(task_id):
    if not require_admin():
        return jsonify({"error": "Unauthorized"}), 401
    data = request.get_json(force=True)

    conn = get_conn()
    cur = conn.cursor()
    cur.execute("SELECT * FROM tasks WHERE id = ?", (task_id,))
    row = cur.fetchone()
    if not row:
        conn.close()
        return jsonify({"error": "Not found"}), 404

    fields, params = [], []
    old_status = row["status"]

    if "status" in data:
        new_status = data["status"]
        if new_status not in ALLOWED_STATUSES or not valid_transition(old_status, new_status):
            conn.close()
            return jsonify({"error": "Invalid status transition"}), 400
        fields.append("status = ?")
        params.append(new_status)
        if new_status == "Done" and not row["completed_at"]:
            fields.append("completed_at = ?"); params.append(now())
        if old_status == "Done" and new_status != "Done":
            fields.append("completed_at = NULL")

    if "progress" in data:
        p = int(data["progress"]); p = max(0, min(100, p))
        fields.append("progress_pct = ?"); params.append(p)

    if "deadline" in data and isinstance(data["deadline"], dict):
        d = data["deadline"]
        typ = d.get("type"); dt = d.get("date")
        if typ not in (None, "soft", "hard"):
            conn.close(); return jsonify({"error": "Invalid deadline type"}), 400
        if dt is not None:
            try: date.fromisoformat(dt)
            except Exception: conn.close(); return jsonify({"error":"Invalid deadline date"}), 400
        fields.append("deadline_type = ?"); params.append(typ)
        fields.append("deadline_date = ?"); params.append(dt)

    if "title" in data:
        title = (data["title"] or "").strip()
        if not title: conn.close(); return jsonify({"error":"Title required"}), 400
        fields.append("title = ?"); params.append(title)

    if "category" in data:
        cat = data["category"]
        if cat not in ALLOWED_CATEGORIES:
            conn.close(); return jsonify({"error":"Invalid category"}), 400
        fields.append("category = ?"); params.append(cat)

    if "tags" in data and isinstance(data["tags"], list):
        cur.execute("DELETE FROM task_tags WHERE task_id = ?", (task_id,))
        for t in data["tags"]:
            tag = (t or "").strip().lower()
            if not tag: continue
            try: cur.execute("INSERT INTO tags(name) VALUES(?)", (tag,))
            except Exception: pass
            cur.execute("SELECT id FROM tags WHERE name = ?", (tag,))
            tag_id = cur.fetchone()["id"]
            cur.execute("INSERT OR IGNORE INTO task_tags(task_id, tag_id) VALUES(?,?)", (task_id, tag_id))

    fields.append("updated_at = ?"); params.append(now())
    params.append(task_id)
    cur.execute(f"UPDATE tasks SET {', '.join(fields)} WHERE id = ?", params)

    conn.commit()
    cur.execute("SELECT * FROM tasks WHERE id = ?", (task_id,))
    task = task_to_dict(cur.fetchone(), conn)
    conn.close()
    return jsonify(task)

@app.delete("/api/tasks/<int:task_id>")
def delete_task(task_id):
    if not require_admin():
        return jsonify({"error": "Unauthorized"}), 401
    conn = get_conn(); cur = conn.cursor()
    cur.execute("DELETE FROM tasks WHERE id = ?", (task_id,))
    conn.commit(); conn.close()
    return jsonify({"ok": True})

# --- Updates ---
@app.post("/api/tasks/<int:task_id>/updates")
def create_update(task_id):
    if not require_admin():
        return jsonify({"error": "Unauthorized"}), 401

    text = request.form.get("text", "").strip()
    if not text:
        return jsonify({"error": "Text required"}), 400

    progress_pct = request.form.get("progress_pct")
    progress_int = None
    if progress_pct not in (None, ""):
        progress_int = max(0, min(100, int(progress_pct)))

    is_highlight = 1 if request.form.get("is_highlight") == "true" else 0

    links_raw = request.form.get("links_json", "[]")
    try:
        links = json.loads(links_raw)
        if not isinstance(links, list): links = []
    except Exception:
        links = []
    links_json = json.dumps(links, ensure_ascii=False)

    image_path = None
    if "image" in request.files:
        image_path = save_image_webp(request.files["image"])

    conn = get_conn(); cur = conn.cursor()
    ts = now()

    cur.execute("SELECT * FROM tasks WHERE id = ?", (task_id,))
    task = cur.fetchone()
    if not task:
        conn.close(); return jsonify({"error":"Task not found"}), 404

    cur.execute("""
      INSERT INTO updates(task_id, created_at, updated_at, text, progress_pct, links_json, image_path, is_highlight)
      VALUES(?,?,?,?,?,?,?,?)
    """, (task_id, ts, ts, text, progress_int, links_json, image_path, is_highlight))
    upd_id = cur.lastrowid

    cur.execute("UPDATE tasks SET last_update_at = ?, updated_at = ? WHERE id = ?", (ts, ts, task_id))
    if progress_int is not None:
        cur.execute("UPDATE tasks SET progress_pct = ?, updated_at = ? WHERE id = ?", (progress_int, ts, task_id))

    conn.commit()
    cur.execute("SELECT * FROM updates WHERE id = ?", (upd_id,))
    upd = dict(cur.fetchone())
    conn.close()
    return jsonify(upd), 201

@app.patch("/api/updates/<int:update_id>")
def edit_update(update_id):
    if not require_admin():
        return jsonify({"error": "Unauthorized"}), 401

    conn = get_conn(); cur = conn.cursor()
    cur.execute("SELECT * FROM updates WHERE id = ?", (update_id,))
    old = cur.fetchone()
    if not old:
        conn.close(); return jsonify({"error":"Not found"}), 404

    text = request.form.get("text", "").strip() if "text" in request.form else old["text"]

    progress_int = old["progress_pct"]
    if "progress_pct" in request.form:
        pv = request.form.get("progress_pct")
        if pv not in (None, ""): progress_int = max(0, min(100, int(pv)))
        else: progress_int = None

    is_highlight = old["is_highlight"]
    if "is_highlight" in request.form:
        is_highlight = 1 if request.form.get("is_highlight") == "true" else 0

    links_json = old["links_json"]
    if "links_json" in request.form:
        try:
            links = json.loads(request.form.get("links_json", "[]"))
            if not isinstance(links, list): links = []
            links_json = json.dumps(links, ensure_ascii=False)
        except Exception:
            pass

    image_path = old["image_path"]
    if "image" in request.files:
        image_path = save_image_webp(request.files["image"])

    ts = now()
    cur.execute("""
      UPDATE updates SET text=?, progress_pct=?, links_json=?, image_path=?, is_highlight=?, updated_at=?
      WHERE id=?
    """, (text, progress_int, links_json, image_path, is_highlight, ts, update_id))

    if progress_int is not None:
        cur.execute("UPDATE tasks SET progress_pct=?, updated_at=? WHERE id=?", (progress_int, ts, old["task_id"]))
    cur.execute("UPDATE tasks SET last_update_at=?, updated_at=? WHERE id=?", (ts, ts, old["task_id"]))

    conn.commit()
    cur.execute("SELECT * FROM updates WHERE id = ?", (update_id,))
    upd = dict(cur.fetchone())
    conn.close()
    return jsonify(upd)

@app.delete("/api/updates/<int:update_id>")
def delete_update(update_id):
    if not require_admin():
        return jsonify({"error": "Unauthorized"}), 401
    conn = get_conn(); cur = conn.cursor()
    cur.execute("SELECT task_id FROM updates WHERE id = ?", (update_id,))
    r = cur.fetchone()
    if not r:
        conn.close(); return jsonify({"ok": True})
    task_id = r["task_id"]
    cur.execute("DELETE FROM updates WHERE id = ?", (update_id,))
    cur.execute("SELECT created_at FROM updates WHERE task_id = ? ORDER BY created_at DESC LIMIT 1", (task_id,))
    r2 = cur.fetchone()
    last = r2["created_at"] if r2 else None
    cur.execute("UPDATE tasks SET last_update_at=?, updated_at=? WHERE id=?", (last, now(), task_id))
    conn.commit(); conn.close()
    return jsonify({"ok": True})

# --- History & Export ---
@app.get("/api/history")
def history():
    page, size, offset = parse_pagination()
    conn = get_conn(); cur = conn.cursor()
    sql = "SELECT * FROM tasks WHERE status='Done'"
    params = []

    frm = request.args.get("from"); to = request.args.get("to")
    if frm: sql += " AND date(completed_at) >= date(?)"; params.append(frm)
    if to:  sql += " AND date(completed_at) <= date(?)"; params.append(to)

    category = request.args.get("category")
    if category in ALLOWED_CATEGORIES:
        sql += " AND category = ?"; params.append(category)

    tag = request.args.get("tag")
    if tag:
        sql += """
          AND id IN (
            SELECT tt.task_id FROM task_tags tt
            JOIN tags tg ON tg.id = tt.tag_id
            WHERE tg.name = LOWER(?)
          )
        """
        params.append(tag)

    q = request.args.get("q")
    if q:
        sql += " AND title LIKE ?"; params.append(f"%{q}%")

    total = conn.execute("SELECT COUNT(1) AS c FROM (" + sql + ")", params).fetchone()["c"]
    sql += " ORDER BY completed_at DESC, id DESC LIMIT ? OFFSET ?"
    params2 = params + [size, offset]
    cur.execute(sql, params2)
    items = [task_to_dict(r, conn) for r in cur.fetchall()]
    conn.close()
    return jsonify({"page": page, "size": size, "total": total, "items": items})

def _query_history_for_export(args):
    conn = get_conn(); cur = conn.cursor()
    sql = "SELECT * FROM tasks WHERE status='Done'"
    params = []
    frm = args.get("from"); to = args.get("to")
    if frm: sql += " AND date(completed_at) >= date(?)"; params.append(frm)
    if to:  sql += " AND date(completed_at) <= date(?)"; params.append(to)
    category = args.get("category")
    if category in ALLOWED_CATEGORIES:
        sql += " AND category = ?"; params.append(category)
    tag = args.get("tag")
    if tag:
        sql += """
          AND id IN (
            SELECT tt.task_id FROM task_tags tt
            JOIN tags tg ON tg.id = tt.tag_id
            WHERE tg.name = LOWER(?)
          )
        """; params.append(tag)
    q = args.get("q")
    if q: sql += " AND title LIKE ?"; params.append(f"%{q}%")
    sql += " ORDER BY completed_at DESC"
    cur.execute(sql, params)
    rows = cur.fetchall()
    items = [task_to_dict(r, conn) for r in rows]
    conn.close()
    return items

@app.get("/api/export.csv")
def export_csv():
    items = _query_history_for_export(request.args)
    import csv
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["ID","Title","Category","Tags","Completed At","Deadline Type","Deadline Date","Progress %"])
    for t in items:
        writer.writerow([
            t["id"], t["title"], t["category"], ", ".join(t["tags"]),
            t.get("completed_at") or "", t["deadline"]["type"] or "", t["deadline"]["date"] or "",
            t.get("progress_pct", 0)
        ])
    resp = Response(output.getvalue(), mimetype="text/csv; charset=utf-8")
    resp.headers["Content-Disposition"] = "attachment; filename=history.csv"
    return resp

@app.get("/api/export.xlsx")
def export_xlsx():
    items = _query_history_for_export(request.args)
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "History"
    headers = ["ID","Title","Category","Tags","Completed At","Deadline Type","Deadline Date","Progress %"]
    ws.append(headers)
    for t in items:
        ws.append([
            t["id"], t["title"], t["category"], ", ".join(t["tags"]),
            t.get("completed_at") or "", t["deadline"]["type"] or "", t["deadline"]["date"] or "",
            t.get("progress_pct", 0)
        ])
    for i in range(1, len(headers)+1):
        ws.column_dimensions[get_column_letter(i)].width = 24
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return Response(
        buf.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=history.xlsx"}
    )

if __name__ == "__main__":
    from waitress import serve  # optional
    try:
        serve(app, host="127.0.0.1", port=5000)
    except Exception:
        app.run(host="127.0.0.1", port=5000, debug=config.DEBUG)
